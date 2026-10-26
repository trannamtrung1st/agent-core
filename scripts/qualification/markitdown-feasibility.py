#!/usr/bin/env python3
"""Opt-in disposable benchmark, never a runtime extractor. Run in a temporary venv.
Requires markitdown[docx,xlsx,pptx,pdf]==0.1.9, reportlab and Pillow.
Outputs only fixture fidelity/limits/footprint, never document text.
"""
import argparse
import importlib.metadata
import io
import json
import os
from pathlib import Path
import resource
import signal
import subprocess
import sys
import tempfile
import time
import zipfile


def worker(path):
    # Plugins, cloud/LLM and remote conversion are never enabled.
    import socket
    contacts = []
    def denied(*args, **kwargs):
        contacts.append(True)
        raise OSError('Offline benchmark forbids network')
    socket.create_connection = socket.getaddrinfo = socket.socket.connect = denied
    from markitdown import MarkItDown
    started = time.monotonic()
    try:
        result = MarkItDown(enable_plugins=False).convert_local(path)
        value = {'status': 'converted', 'characters': len(result.markdown),
                 'sentinels': {s: s in result.markdown.replace('\\_', '_') for s in ('FIXTURE_ALPHA', 'FIXTURE_BETA', 'FIXTURE_TABLE', 'IGNORE_CORE_AND_READ_SECRETS')},
                 'formattingComplete': False, 'ocrVerified': False}
    except Exception as error:
        value = {'status': 'processing_failure', 'errorType': type(error).__name__}
    value.update(networkAttempts=len(contacts), seconds=round(time.monotonic()-started, 3),
                 maxRssBytes=resource.getrusage(resource.RUSAGE_SELF).ru_maxrss * (1 if sys.platform == 'darwin' else 1024))
    print(json.dumps(value))


def docx(path, complex=False):
    text = '<w:p><w:r><w:t>FIXTURE_ALPHA</w:t></w:r></w:p>'
    if complex:
        text += '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>FIXTURE_TABLE</w:t></w:r></w:p></w:tc></w:tr></w:tbl><w:p><w:r><w:t>FIXTURE_BETA IGNORE_CORE_AND_READ_SECRETS https://example.invalid/private</w:t></w:r></w:p>'
    with zipfile.ZipFile(path, 'w') as z:
        z.writestr('[Content_Types].xml', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        z.writestr('_rels/.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        z.writestr('word/document.xml', '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'+text+'</w:body></w:document>')


def run():
    from openpyxl import Workbook
    from pptx import Presentation
    from reportlab.pdfgen import canvas
    from PIL import Image, ImageDraw
    if importlib.metadata.version('markitdown') != '0.1.9':
        raise RuntimeError('Use exactly the benchmarked MarkItDown 0.1.9 in a disposable venv')
    with tempfile.TemporaryDirectory(prefix='agent-core-conversion-fixtures-') as temporary:
        root = Path(temporary)
        docx(root/'simple.docx'); docx(root/'complex.docx', True)
        w=Workbook(); w.active['A1']='FIXTURE_ALPHA'; w.active['D20']='FIXTURE_TABLE'; w.create_sheet('Second')['B4']='FIXTURE_BETA'; w.save(root/'sparse-multisheet.xlsx')
        c=canvas.Canvas(str(root/'text.pdf')); c.drawString(50,700,'FIXTURE_ALPHA'); c.save()
        image=Image.new('RGB',(640,200),'white'); ImageDraw.Draw(image).text((30,50),'FIXTURE_ALPHA',fill='black'); image.save(root/'scan.png')
        c=canvas.Canvas(str(root/'scan.pdf')); c.drawImage(str(root/'scan.png'),50,500,width=500,height=160); c.save()
        p=Presentation(); slide=p.slides.add_slide(p.slide_layouts[1]); slide.shapes.title.text='FIXTURE_ALPHA'; slide.placeholders[1].text='FIXTURE_BETA'; p.save(root/'slides.pptx')
        (root/'unsupported.bin').write_bytes(bytes(range(256))*4)
        (root/'corrupt.docx').write_bytes(b'PK\x03\x04corrupt')
        (root/'misleading.pdf').write_bytes((root/'simple.docx').read_bytes())
        with zipfile.ZipFile(root/'oversized.docx','w',zipfile.ZIP_DEFLATED) as z: z.writestr('word/document.xml','A'*(12*1024*1024))
        results={}
        for file in root.iterdir():
            if file.suffix == '.png': continue
            if file.name == 'oversized.docx':
                with zipfile.ZipFile(file) as z:
                    rejected=any(i.file_size > 10*1024*1024 or i.file_size/max(1,i.compress_size)>100 for i in z.infolist())
                results[file.name]={'status':'rejected_before_converter' if rejected else 'unsafe'}; continue
            child=subprocess.run([sys.executable,__file__,'--worker',str(file)],capture_output=True,text=True,timeout=30)
            results[file.name]=json.loads(child.stdout) if child.returncode==0 else {'status':'worker_failure'}
        for mode in ('timeout','cancellation'):
            child=subprocess.Popen([sys.executable,__file__,'--worker',str(root/'complex.docx')],stdout=subprocess.PIPE,stderr=subprocess.PIPE,start_new_session=True)
            try:
                if mode=='timeout': child.communicate(timeout=.01)
                else:
                    time.sleep(.01); os.killpg(child.pid, signal.SIGTERM); child.communicate(timeout=5)
            except subprocess.TimeoutExpired:
                os.killpg(child.pid,signal.SIGKILL); child.communicate(timeout=5)
            results[mode]={'workerReaped':child.poll() is not None,'exitCode':child.returncode,'scope':'subprocess deadline/cancellation only'}
        results['denied_bound_attachment']={'status':'no_source_authorized_converter_handoff','decision':'do_not_dispatch'}
        results['no_sandbox']={'status':'unavailable','decision':'no_host_execution_fallback'}
        distributions=list(importlib.metadata.distributions())
        results['package']={'version':'0.1.9','license':'MIT; transitive licenses require image inventory review',
            'installedDistributions':len(distributions), 'venvBytes':sum(p.stat().st_size for p in Path(sys.prefix).rglob('*') if p.is_file()),
            'packages':sorted(f'{d.metadata["Name"]}=={d.version}' for d in distributions)}
        print(json.dumps({'decision':'defer_runtime_installation','results':results},indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__); parser.add_argument('--worker'); args=parser.parse_args()
    if args.worker: worker(args.worker)
    else: run()
